# -*- coding: utf-8 -*-
"""docx 大纲 + 3D打印PDF全文提取"""
import sys, io, re
from pathlib import Path
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

ROOT = Path(__file__).resolve().parents[2] / "_extract"
OUT = ROOT / "_read"
OUT.mkdir(exist_ok=True)

from docx import Document
doc = Document(ROOT / "原辅材料技术标准.docx")

# ---- docx 大纲：标题样式段落 + 表标题附近 ----
lines = []
for i, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    if not t:
        continue
    st = (p.style.name if p.style else "")
    is_heading = ("head" in st.lower()) or ("标题" in st)
    # 数字编号开头的小节标题：如 "1 2" "3.2" "5 孕育剂"
    num_head = re.match(r"^\d+(\.\d+)?\s+\S", t) and len(t) < 40
    if is_heading or num_head:
        lines.append(f"[{i}] <{st}> {t}")
    elif re.match(r"^\d+\.\d+", t) and len(t) < 30:
        lines.append(f"[{i}] <{st}> {t}")
with open(OUT / "docx_outline.txt", "w", encoding="utf-8") as f:
    f.write("\n".join(lines))
print(f"docx 大纲 {len(lines)} 条 -> docx_outline.txt")

# ---- 表格索引 ----
with open(OUT / "docx_tables_index.txt", "w", encoding="utf-8") as f:
    for ti, t in enumerate(doc.tables):
        # 取表格首行前3格作为索引
        try:
            head = " | ".join(c.text.strip().replace("\n", " ")[:20] for c in t.rows[0].cells[:3])
        except Exception:
            head = "?"
        f.write(f"表{ti}: {t.rows.__len__()}行x{len(t.columns)}列 | {head}\n")
print(f"docx 表索引 -> docx_tables_index.txt")

# ---- 3D打印PDF 全文 ----
import fitz
jobs = [
    ("问题知识库", "3D打印机故障问题解决知识库.pdf"),
    ("问题知识库", "3D打印砂芯质量问题知识库.pdf"),
    ("问题知识库", "再生砂系统系统故障措施知识库.pdf"),
    ("设备维护保养", "设备三级保养评分表.pdf"),
    ("设备维护保养", "铸铁事业部-刘佃柱-铸造3D打印机日常维护保养操作.pdf"),
]
for folder, name in jobs:
    d = fitz.open(ROOT / folder / name)
    txt = []
    for pi, pg in enumerate(d):
        txt.append(f"\n===== 第{pi+1}页 =====\n")
        t = pg.get_text().strip()
        if t:
            txt.append(t)
        else:
            # 无文本层的页标记为图片页
            imgs = pg.get_images()
            txt.append(f"[图片页, 内嵌图片 {len(imgs)} 张]")
    fn = OUT / (name.replace(".pdf", "") + ".txt")
    fn.write_text("".join(txt), encoding="utf-8")
    print(f"{name}: {len(''.join(txt))}字符 -> {fn.name}")
    d.close()
