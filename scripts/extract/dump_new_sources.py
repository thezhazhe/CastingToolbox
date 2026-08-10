# -*- coding: utf-8 -*-
# 抽取两个新资料 → 文本（避免控制台中文乱码，直接写文件）
#   1. 国内外材料对照表.xls        → _extract/_read/材料对照表.txt
#   2. 铸造铝合金3D打印砂型工艺设计规范.docx → _extract/_read/铝合金3DP规范.txt
import os, sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

ROOT = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))
SRC = r'D:\技术文件\学习资料'
OUT = os.path.join(ROOT, 'scripts', 'extract', '_extract', '_read')
os.makedirs(OUT, exist_ok=True)

def main():
    # ---- 1. XLS 材料对照表 ----
    xls_path = os.path.join(SRC, '国内外材料对照表.xls')
    import xlrd
    book = xlrd.open_workbook(xls_path)
    lines = []
    for sh in book.sheets():
        lines.append(f'=== SHEET: {sh.name} ({sh.nrows}x{sh.ncols}) ===')
        for r in range(sh.nrows):
            row = []
            for c in range(sh.ncols):
                v = sh.cell_value(r, c)
                if isinstance(v, float) and v.is_integer():
                    v = int(v)
                row.append('' if v == '' else str(v))
            lines.append(' | '.join(row))
    with open(os.path.join(OUT, '材料对照表.txt'), 'w', encoding='utf-8') as f:
        f.write('\n'.join(lines))
    print(f'XLS: {len(book.sheets())} sheets dumped')

    # ---- 2. DOCX 铝合金3D打印砂型工艺设计规范 ----
    docx_path = os.path.join(SRC, '铸造铝合金3D打印砂型工艺设计规范-20191226_(1).docx')
    import docx
    doc = docx.Document(docx_path)
    lines2 = []
    # 段落
    for p in doc.paragraphs:
        t = p.text.strip()
        if t:
            lines2.append(t)
    # 表格
    for ti, table in enumerate(doc.tables):
        lines2.append(f'--- TABLE {ti+1} ---')
        for row in table.rows:
            cells = [c.text.strip().replace('\n', ' ') for c in row.cells]
            lines2.append(' | '.join(cells))
    with open(os.path.join(OUT, '铝合金3DP规范.txt'), 'w', encoding='utf-8') as f:
        f.write('\n'.join(lines2))
    print(f'DOCX: {len(doc.paragraphs)} paragraphs, {len(doc.tables)} tables dumped')

if __name__ == '__main__':
    main()
