# -*- coding: utf-8 -*-
"""抽查 docx：表1 物料清单 + 表2 + 表73 钼铁 + 开头段落"""
import sys, io
from pathlib import Path
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

ROOT = Path(__file__).resolve().parents[2] / "_extract"
from docx import Document
doc = Document(ROOT / "原辅材料技术标准.docx")

def dump_table(ti, maxrows=100, maxcols=12):
    t = doc.tables[ti]
    print(f"\n===== 表{ti} ({len(t.rows)}行 x {len(t.columns)}列) =====")
    for ri, r in enumerate(t.rows):
        if ri >= maxrows:
            print(f"...(共{len(t.rows)}行)")
            break
        cells = []
        for c in r.cells[:maxcols]:
            txt = c.text.strip().replace("\n", "/")
            cells.append(txt[:40])
        print(f"[{ri}] " + " | ".join(cells))

print("############ 表0 更改记录（检查脱敏项）############")
dump_table(0, maxrows=10)
print("\n############ 表1 物料总清单 ############")
dump_table(1, maxrows=130, maxcols=8)

print("\n############ 表2 第一种材料规范 ############")
dump_table(2, maxrows=20)

print("\n############ 表73 钼铁 ############")
dump_table(73, maxrows=20)

print("\n############ 开头段落 0~80 ############")
for i, p in enumerate(doc.paragraphs[:80]):
    if p.text.strip():
        print(f"[{i}] {p.text.strip()[:120]}")
