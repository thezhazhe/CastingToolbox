# -*- coding: utf-8 -*-
"""抽查配料单 xlsx 结构"""
import sys, io
from pathlib import Path
import openpyxl
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

DIR = Path(r"D:\CCproject\CastingToolbox\_extract\配料单")
files = ["QT450-10 (2).xlsx", "中硅钼.xlsx", "6411制动底板.xlsx", "缸体.xlsx"]
for fn in files:
    wb = openpyxl.load_workbook(DIR / fn, data_only=True)
    print(f"\n############ {fn} 工作表: {wb.sheetnames} ############")
    ws = wb[wb.sheetnames[0]]
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 30)):
        vals = []
        for c in row:
            if c.value is not None:
                vals.append(f"{c.coordinate}={str(c.value)[:28]}")
        if vals:
            print(" | ".join(vals))
