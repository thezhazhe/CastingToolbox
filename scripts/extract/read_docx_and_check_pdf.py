# -*- coding: utf-8 -*-
"""读取 docx 全文 + 检查各 PDF 文本层情况，输出到 _extract/_read/"""
import sys, io
from pathlib import Path
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

ROOT = Path(__file__).resolve().parents[2] / "_extract"
OUT = ROOT / "_read"
OUT.mkdir(exist_ok=True)

# ---- 1. docx 全文 ----
from docx import Document
doc = Document(ROOT / "原辅材料技术标准.docx")
paras = [p.text for p in doc.paragraphs]
# 表格
tables = []
for ti, t in enumerate(doc.tables):
    rows = []
    for r in t.rows:
        rows.append([c.text.strip() for c in r.cells])
    tables.append((ti, rows))

with open(OUT / "docx_text.txt", "w", encoding="utf-8") as f:
    f.write("===== 段落 =====\n")
    for i, p in enumerate(paras):
        if p.strip():
            f.write(f"[{i}] {p}\n")
    f.write("\n===== 表格 =====\n")
    for ti, rows in tables:
        f.write(f"\n--- 表{ti} ({len(rows)}行) ---\n")
        for r in rows[:80]:
            f.write(" | ".join(r) + "\n")
print(f"docx: {len(paras)}段, {len(tables)}表 -> docx_text.txt")

# ---- 2. PDF 文本层检查 ----
import fitz
pdfs = [
    ROOT / "问题知识库" / "3D打印机故障问题解决知识库.pdf",
    ROOT / "问题知识库" / "3D打印砂芯质量问题知识库.pdf",
    ROOT / "问题知识库" / "再生砂系统系统故障措施知识库.pdf",
    ROOT / "设备维护保养" / "设备三级保养评分表.pdf",
    ROOT / "设备维护保养" / "铸铁事业部-刘佃柱-铸造3D打印机日常维护保养操作.pdf",
]
for p in pdfs:
    d = fitz.open(p)
    pages = d.page_count
    total_chars = 0
    for pg in d:
        total_chars += len(pg.get_text().strip())
    imgs = sum(len(d[i].get_images()) for i in range(pages))
    print(f"{p.name}: {pages}页, 文本层字符={total_chars}, 内嵌图片={imgs}")
    d.close()
